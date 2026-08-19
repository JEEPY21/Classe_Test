"""Text extraction from uploaded CV files (PDF / DOCX).
The extracted text is cleaned before being sent to the LLM: collapse repeated whitespace,
drop empty/duplicate blank lines, and strip repeated headers/footers across pages."""
import io
import re
from collections import Counter
import pdfplumber
from docx import Document


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)              # collapse multiple spaces/tabs
    lines = [ln.strip() for ln in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)                 # collapse excessive blank lines
    return text.strip()


def _strip_repeated_headers_footers(pages: list) -> str:
    """Remove short lines that repeat across multiple pages (typical headers/footers such as a
    name band, page numbers or a URL). Only applied when the document has 2+ pages."""
    page_lines = [[ln.strip() for ln in (p or "").split("\n")] for p in pages]
    repeated = set()
    if len(pages) >= 2:
        counts = Counter()
        for pl in page_lines:
            for ln in {x for x in pl if x}:
                counts[ln] += 1
        for ln, c in counts.items():
            # short line appearing on at least 2 pages -> header/footer artefact
            if c >= 2 and len(ln) <= 80:
                repeated.add(ln)
            # page-number style lines
            if re.fullmatch(r"(page\s*)?\d+(\s*/\s*\d+)?", ln, flags=re.IGNORECASE):
                repeated.add(ln)
    kept = []
    for pl in page_lines:
        for ln in pl:
            if ln in repeated:
                continue
            kept.append(ln)
    return "\n".join(kept)


def extract_text_from_pdf(data: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = _strip_repeated_headers_footers(pages)
    return _normalize_whitespace(text)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalize_whitespace("\n".join(parts))


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    if name.endswith(".docx"):
        return extract_text_from_docx(data)
    raise ValueError("Format non supporté. Utilisez PDF ou DOCX.")
