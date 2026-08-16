"""BLOC 5 — Layout & export. The LLM never touches layout.
Fixed, ATS-friendly HTML/CSS templates (single column, standard section names,
no tables, no multi-column key info). Native-text PDF (xhtml2pdf) + DOCX (python-docx)."""
import io
import html
from xhtml2pdf import pisa
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SECTION_LABELS = {
    "fr": {"summary": "Profil", "experience": "Expérience professionnelle",
           "education": "Formation", "skills": "Compétences", "languages": "Langues"},
    "en": {"summary": "Summary", "experience": "Professional Experience",
           "education": "Education", "skills": "Skills", "languages": "Languages"},
    "es": {"summary": "Perfil", "experience": "Experiencia profesional",
           "education": "Formación", "skills": "Competencias", "languages": "Idiomas"},
    "de": {"summary": "Profil", "experience": "Berufserfahrung",
           "education": "Ausbildung", "skills": "Kompetenzen", "languages": "Sprachen"},
}

TEMPLATES = [
    {"id": "classic", "name": "Classique", "accent": "#0F172A", "font": "Helvetica",
     "description": "Sobre et intemporel, une colonne, titres soulignés."},
    {"id": "modern", "name": "Moderne", "accent": "#2563EB", "font": "Helvetica",
     "description": "Accent bleu, titres en majuscules espacées, aéré."},
    {"id": "compact", "name": "Compact", "accent": "#334155", "font": "Times",
     "description": "Dense et efficace, idéal pour les parcours riches."},
]


def build_final_cv(cv: dict, session: dict) -> dict:
    """Assemble the reviewed final CV from base CV + accepted rewrites + confirmed skills."""
    decisions = session.get("decisions", {})
    rewrite_map = {r["experience_id"]: r for r in session.get("experience_rewrites", [])}
    accepted_map = {d["experience_id"]: set(d.get("accepted_bullets", []))
                    for d in decisions.get("rewrites", [])}

    final = {
        "identity": cv.get("identity", {}),
        "summary": cv.get("summary", ""),
        "education": cv.get("education", []),
        "languages": cv.get("languages", []),
        "experiences": [],
        "skills": list(cv.get("skills", [])),
    }
    for exp in cv.get("experiences", []):
        eid = exp.get("id")
        rewrite = rewrite_map.get(eid)
        accepted = accepted_map.get(eid, set())
        bullets = []
        if rewrite:
            for i, b in enumerate(rewrite.get("bullets", [])):
                bullets.append(b["proposed"] if i in accepted else b.get("original", ""))
        else:
            bullets = exp.get("bullets", [])
        bullets = [b for b in bullets if b and b.strip()]
        final["experiences"].append({**exp, "bullets": bullets})

    for s in session.get("confirmed_skills", []):
        if s not in final["skills"]:
            final["skills"].append(s)
    return final


def _e(v) -> str:
    return html.escape(str(v or ""))


def _date_range(start, end, current):
    parts = []
    if start:
        parts.append(_e(start))
    if current:
        parts.append("présent")
    elif end:
        parts.append(_e(end))
    return " – ".join(parts)


def render_html(cv: dict, template_id: str, language: str = "fr") -> str:
    tpl = next((t for t in TEMPLATES if t["id"] == template_id), TEMPLATES[0])
    labels = SECTION_LABELS.get(language, SECTION_LABELS["fr"])
    accent = tpl["accent"]
    font = tpl["font"]
    upper = "text-transform:uppercase;letter-spacing:2px;" if tpl["id"] == "modern" else ""
    fs = "10.5px" if tpl["id"] == "compact" else "11.5px"
    ident = cv.get("identity", {})

    contact = " · ".join([x for x in [ident.get("email"), ident.get("phone"),
                                      ident.get("location"), *(ident.get("links") or [])] if x])

    def section_title(t):
        return (f'<div style="font-size:13px;font-weight:bold;color:{accent};{upper}'
                f'border-bottom:1.5px solid {accent};padding-bottom:3px;margin:16px 0 8px 0;">{_e(t)}</div>')

    parts = [f"""<html><head><meta charset="utf-8"><style>
      @page {{ size: A4; margin: 1.6cm; }}
      body {{ font-family: {font}; font-size:{fs}; color:#1a1a1a; line-height:1.5; }}
      h1 {{ font-size:22px; color:{accent}; margin:0 0 2px 0; }}
      .title {{ font-size:13px; color:#444; margin:0 0 4px 0; }}
      .contact {{ font-size:10px; color:#555; margin-bottom:6px; }}
      .exp-head {{ font-weight:bold; font-size:12px; }}
      .exp-sub {{ font-style:italic; color:#555; margin-bottom:2px; }}
      ul {{ margin:4px 0 8px 0; padding-left:16px; }}
      li {{ margin-bottom:2px; }}
      .skills {{ margin:0; }}
    </style></head><body>"""]

    parts.append(f'<h1>{_e(ident.get("full_name"))}</h1>')
    if ident.get("title"):
        parts.append(f'<div class="title">{_e(ident.get("title"))}</div>')
    if contact:
        parts.append(f'<div class="contact">{contact}</div>')

    if cv.get("summary"):
        parts.append(section_title(labels["summary"]))
        parts.append(f'<div>{_e(cv["summary"])}</div>')

    if cv.get("experiences"):
        parts.append(section_title(labels["experience"]))
        for exp in cv["experiences"]:
            dr = _date_range(exp.get("start_date"), exp.get("end_date"), exp.get("current"))
            head = _e(exp.get("role"))
            if exp.get("company"):
                head += f' — {_e(exp.get("company"))}'
            parts.append(f'<div class="exp-head">{head}</div>')
            sub = " · ".join([x for x in [dr, _e(exp.get("location")) if exp.get("location") else ""] if x])
            if sub:
                parts.append(f'<div class="exp-sub">{sub}</div>')
            if exp.get("bullets"):
                parts.append("<ul>" + "".join(f"<li>{_e(b)}</li>" for b in exp["bullets"]) + "</ul>")

    if cv.get("education"):
        parts.append(section_title(labels["education"]))
        for ed in cv["education"]:
            dr = _date_range(ed.get("start_date"), ed.get("end_date"), False)
            line = _e(ed.get("degree"))
            if ed.get("field"):
                line += f' — {_e(ed.get("field"))}'
            parts.append(f'<div class="exp-head">{line}</div>')
            sub = " · ".join([x for x in [_e(ed.get("institution")), dr] if x])
            if sub:
                parts.append(f'<div class="exp-sub">{sub}</div>')

    if cv.get("skills"):
        parts.append(section_title(labels["skills"]))
        parts.append(f'<div class="skills">{" · ".join(_e(s) for s in cv["skills"])}</div>')

    if cv.get("languages"):
        parts.append(section_title(labels["languages"]))
        langs = [f'{_e(l.get("language"))} ({_e(l.get("level"))})' if l.get("level") else _e(l.get("language"))
                 for l in cv["languages"]]
        parts.append(f'<div>{" · ".join(langs)}</div>')

    parts.append("</body></html>")
    return "".join(parts)


def html_to_pdf(html_str: str) -> bytes:
    buf = io.BytesIO()
    pisa.CreatePDF(src=html_str, dest=buf, encoding="utf-8")
    return buf.getvalue()


def build_docx(cv: dict, language: str = "fr") -> bytes:
    labels = SECTION_LABELS.get(language, SECTION_LABELS["fr"])
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    ident = cv.get("identity", {})

    h = doc.add_paragraph()
    run = h.add_run(ident.get("full_name", ""))
    run.bold = True
    run.font.size = Pt(18)
    if ident.get("title"):
        doc.add_paragraph(ident["title"])
    contact = " · ".join([x for x in [ident.get("email"), ident.get("phone"),
                                      ident.get("location"), *(ident.get("links") or [])] if x])
    if contact:
        doc.add_paragraph(contact)

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if cv.get("summary"):
        heading(labels["summary"])
        doc.add_paragraph(cv["summary"])

    if cv.get("experiences"):
        heading(labels["experience"])
        for exp in cv["experiences"]:
            p = doc.add_paragraph()
            title = exp.get("role", "")
            if exp.get("company"):
                title += f' — {exp["company"]}'
            p.add_run(title).bold = True
            dr = _date_range(exp.get("start_date"), exp.get("end_date"), exp.get("current")).replace("&#x2013;", "–")
            sub = " · ".join([x for x in [dr, exp.get("location", "")] if x])
            if sub:
                doc.add_paragraph(sub).runs[0].italic = True
            for b in exp.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")

    if cv.get("education"):
        heading(labels["education"])
        for ed in cv["education"]:
            p = doc.add_paragraph()
            line = ed.get("degree", "")
            if ed.get("field"):
                line += f' — {ed["field"]}'
            p.add_run(line).bold = True
            sub = " · ".join([x for x in [ed.get("institution", ""),
                                          _date_range(ed.get("start_date"), ed.get("end_date"), False)] if x])
            if sub:
                doc.add_paragraph(sub)

    if cv.get("skills"):
        heading(labels["skills"])
        doc.add_paragraph(" · ".join(cv["skills"]))

    if cv.get("languages"):
        heading(labels["languages"])
        langs = [f'{l.get("language","")} ({l.get("level","")})' if l.get("level") else l.get("language", "")
                 for l in cv["languages"]]
        doc.add_paragraph(" · ".join(langs))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
